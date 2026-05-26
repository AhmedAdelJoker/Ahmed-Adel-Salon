from fastapi.middleware.cors import CORSMiddleware
from datetime import date
from io import BytesIO

from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.api.deps import require_owner
from app.db.session import get_db
from app.models.activity_log import ActivityLog
from app.models.user import User
from app.services.activity_service import log_activity
from app.services.runtime_settings_service import get_runtime_settings

router = APIRouter(prefix="/exports", tags=["Runtime Exports"])

EXCEL_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

DEFAULT_SECURITY_SETTINGS = {
    "enforceStrongPasswords": True,
    "requireShiftForSales": True,
    "lockClosedShiftEdits": True,
    "enableActivityLogs": True,
    "restrictExportsToManagers": True,
    "requireDiscountApproval": True,
    "sessionTimeoutMinutes": 60,
    "maxFailedLoginAttempts": 5,
}

DEFAULT_FINANCIAL_RULES = {
    "taxRate": 0,
    "cashierDiscountLimit": 50,
    "requireDiscountReason": True,
    "requireManagerApprovalAboveLimit": True,
    "requireOpenShiftForInvoices": True,
    "lockInvoicesAfterShiftClose": True,
    "allowNegativeInventory": False,
    "enableProductCostTracking": True,
    "enableExpenseApproval": True,
    "maxCashDiscrepancyWithoutNote": 0,
    "defaultPaymentMethod": "cash",
    "enabledPaymentMethods": {
        "cash": True,
        "card": True,
        "wallet": True,
        "instapay": True,
    },
}


def _workbook_response(workbook: Workbook, filename: str) -> StreamingResponse:
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type=EXCEL_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _apply_header_style(worksheet, columns: list[str]) -> None:
    header_fill = PatternFill("solid", fgColor="0EA5E9")
    header_font = Font(bold=True, color="FFFFFF")
    right_align = Alignment(horizontal="right", vertical="center")

    worksheet.append(columns)
    worksheet.sheet_view.rightToLeft = True
    worksheet.freeze_panes = "A2"
    for index, _ in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = right_align


def _autosize(worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
            cell.alignment = Alignment(horizontal="right", vertical="center")
        worksheet.column_dimensions[column_letter].width = min(
            max(max_length + 3, 18), 40
        )


@router.get("/security-audit")
def export_security_audit(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_owner),
):
    workbook = Workbook()
    settings_sheet = workbook.active
    settings_sheet.title = "Security Settings"
    _apply_header_style(settings_sheet, ["الإعداد", "القيمة"])

    settings_data = get_runtime_settings("security_settings", DEFAULT_SECURITY_SETTINGS)
    for key, value in settings_data.items():
        settings_sheet.append([key, str(value)])

    audit_sheet = workbook.create_sheet("Security Audit")
    _apply_header_style(
        audit_sheet,
        ["الإجراء", "الكيان", "الوصف", "المستخدم", "التاريخ"],
    )

    logs = (
        db.query(ActivityLog)
        .order_by(ActivityLog.created_at.desc())
        .limit(100)
        .all()
    )
    for row in logs:
        audit_sheet.append(
            [
                row.action,
                row.entity_type,
                row.description or "",
                getattr(row.user, "username", "") if row.user else "",
                row.created_at.isoformat() if row.created_at else "",
            ]
        )

    _autosize(settings_sheet)
    _autosize(audit_sheet)
    log_activity(
        db,
        user_id=current_user.id,
        action="export_security_audit",
        entity_type="security_settings",
        description="تم تصدير تقرير الأمان",
    )
    return _workbook_response(
        workbook,
        f"security_audit_{date.today().isoformat()}.xlsx",
    )


@router.get("/financial-rules")
def export_financial_rules(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_owner),
):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Financial Rules"
    _apply_header_style(sheet, ["الإعداد", "القيمة"])

    rules_data = get_runtime_settings("financial_rules", DEFAULT_FINANCIAL_RULES)
    for key, value in rules_data.items():
        sheet.append([key, str(value)])

    _autosize(sheet)
    log_activity(
        db,
        user_id=current_user.id,
        action="export_financial_rules",
        entity_type="financial_rules",
        description="تم تصدير القواعد المالية",
    )
    return _workbook_response(
        workbook,
        f"financial_rules_{date.today().isoformat()}.xlsx",
    )



